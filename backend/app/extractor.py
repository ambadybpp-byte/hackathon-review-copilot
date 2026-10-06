from pathlib import Path
from pptx import Presentation
from pypdf import PdfReader


def extract_pptx(path: str) -> str:
    presentation = Presentation(path)

    chunks = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):
        slide_text = []

        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_text.append(shape.text.strip())

        if slide_text:
            chunks.append(
                f"[SLIDE {slide_number}]\n"
                + "\n".join(slide_text)
            )

    return "\n\n".join(chunks)


def extract_pdf(path: str) -> str:
    reader = PdfReader(path)

    chunks = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""

        if text.strip():
            chunks.append(
                f"[PAGE {page_number}]\n{text.strip()}"
            )

    return "\n\n".join(chunks)


def extract_document(path: str) -> str:
    extension = Path(path).suffix.lower()

    if extension == ".pptx":
        return extract_pptx(path)

    if extension == ".pdf":
        return extract_pdf(path)

    raise ValueError(
        f"Unsupported document type: {extension}"
    )